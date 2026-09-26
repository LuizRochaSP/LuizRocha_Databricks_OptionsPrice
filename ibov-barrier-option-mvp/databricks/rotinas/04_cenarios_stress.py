# Databricks notebook source
# /// script
# [tool.databricks.environment]
# environment_version = "5"
# dependencies = [
#   "openpyxl>=3.1",
#   "exchange-calendars==4.13.2",
#   "lxml>=6.0",
# ]
# ///
# DBTITLE 1,Título
# MAGIC %md
# MAGIC # 04 — Cenários e Stress Test
# MAGIC
# MAGIC Consome um conjunto coerente de RESULTS, VALIDATION e RISK aprovados,
# MAGIC lê cenários de stress de uma planilha Excel, reprecifica a opção com
# MAGIC choques nos fatores de mercado e calcula o P&L da posição em pontos do índice.

# COMMAND ----------

# DBTITLE 1,Parâmetros de configuração
# Comentário: caminho do Excel de entrada com os cenários.
STRESS_EXCEL_PATH = ""  # vazio = usar template padrão em templates/stress_scenarios.xlsx

# COMMAND ----------

# DBTITLE 1,Dependências e leitura dos arquivos de entrada
# Comentário: localiza a raiz do projeto e importa o motor de stress.
import hashlib
import io
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd

def find_project_root() -> Path:
    for candidate in [Path.cwd(), *Path.cwd().parents]:
        if (candidate / "data").exists() and (candidate / "src").exists():
            return candidate
    raise FileNotFoundError("Não encontrei a raiz do projeto contendo data e src.")

project_root = find_project_root()
src_path = project_root / "src"
if str(src_path) not in sys.path:
    sys.path.insert(0, str(src_path))

results_root = project_root.parents[1] / "ibov-barrier-results"

# --- Localiza VALIDATION mais recente ---
validation_pattern = re.compile(r"^VALIDATION_superficieVol_(\d{8})_(\d{6})\.json$")
validation_candidates = []
for path in results_root.glob("VALIDATION_superficieVol_*.json"):
    match = validation_pattern.match(path.name)
    if match:
        stamp = datetime.strptime("".join(match.groups()), "%Y%m%d%H%M%S")
        validation_candidates.append((stamp, path))

if not validation_candidates:
    raise RuntimeError("Nenhuma VALIDATION encontrada. Execute os notebooks 01 e 02 primeiro.")

validation_stamp, validation_path = max(validation_candidates, key=lambda item: item[0])
validation = json.loads(validation_path.read_text(encoding="utf-8"))

# --- Localiza RESULTS referenciado pela VALIDATION ---
results_path = results_root / validation["results_file"]
if not results_path.is_file():
    raise RuntimeError(f"RESULTS não encontrado: {results_path}")

# --- Localiza RISK correspondente (opcional) ---
risk_filename = validation_path.name.replace("VALIDATION_", "RISK_")
risk_path = results_root / risk_filename
risk = json.loads(risk_path.read_text(encoding="utf-8")) if risk_path.is_file() else None

# --- Verifica proveniência ---
from ibov_barrier.stress import (
    ScenarioInput,
    SimulationParams,
    run_all_scenarios,
    verify_provenance,
)

provenance = verify_provenance(results_path, validation, risk)

metadata = provenance["metadata"]
pricing_result = provenance["pricing_result"]
simulation = provenance["simulation"]
results_sha256 = provenance["results_sha256"]

print(f"✓ Validação: {validation_path.name}")
print(f"✓ RESULTS: {results_path.name}")
print(f"✓ SHA-256: {results_sha256}")
print(f"✓ Decisão: {validation['decision']} | WARN={validation['warn_count']} | FAIL={validation['fail_count']}")
if risk:
    print(f"✓ RISK: {risk_path.name}")
print(f"✓ Data de mercado: {metadata['market_date']}")

# COMMAND ----------

# DBTITLE 1,Reconstrói mercado, contrato e parâmetros de simulação
from ibov_barrier import BarrierContract, MarketData

market_values = pricing_result["market"]
contract_values = pricing_result["contract"]

base_market = MarketData(
    spot=float(market_values["spot"]),
    rate=float(market_values["rate"]),
    dividend_yield=float(market_values["dividend_yield"]),
    volatility=float(market_values["volatility"]),
)
base_contract = BarrierContract(
    strike=float(contract_values["strike"]),
    barrier=float(contract_values["barrier"]),
    maturity=float(contract_values["maturity"]),
)
sim = SimulationParams(
    paths=int(simulation["paths"]),
    steps=int(simulation["steps"]),
    seed=int(simulation["seed"]),
    monitoring=str(simulation["monitoring"]),
)

print(f"✓ Mercado base: spot={base_market.spot:,.2f}, r={base_market.rate:.4%}, q={base_market.dividend_yield:.4%}, σ={base_market.volatility:.4%}")
print(f"✓ Contrato: K={base_contract.strike:,.2f}, H={base_contract.barrier:,.2f}, T={base_contract.maturity:.4f}")
print(f"✓ Simulação: {sim.paths:,} caminhos × {sim.steps} passos, seed={sim.seed}, monitoring={sim.monitoring}")

# COMMAND ----------

# DBTITLE 1,Leitura dos cenários do Excel
# Comentário: lê cada aba do Excel como um cenário independente.
# Layout da aba:
#   Linha 1: nome do cenário (título)
#   Linha 2: spot_shock_pct (ex.: -0.10 = -10%)
#   Linha 3: vol_shock_abs (ex.: +0.05 = +5pp)
#   Linha 4: rate_shock_bp (ex.: +100 = +100bp)
#   Linha 5: lado da posição — COMPRADO ou VENDIDO
#   Linha 6: QtdPontos
#   A partir da linha 10: respostas (preenchidas pelo notebook)

if not STRESS_EXCEL_PATH:
    excel_path = project_root / "templates" / "stress_scenarios.xlsx"
else:
    excel_path = Path(STRESS_EXCEL_PATH)

if not excel_path.is_file():
    raise RuntimeError(f"Excel de cenários não encontrado: {excel_path}")

# Preserva o arquivo de entrada — lê uma cópia em memória
input_bytes = excel_path.read_bytes()
xls = pd.ExcelFile(io.BytesIO(input_bytes), engine="openpyxl")

scenarios: list[ScenarioInput] = []
sheet_names = [s for s in xls.sheet_names if s != "RESUMO"]

for sheet in sheet_names:
    df = xls.parse(sheet, header=None, engine="openpyxl")
    try:
        name = str(df.iloc[0, 1]).strip() if df.shape[0] > 0 and pd.notna(df.iloc[0, 1]) else sheet
        spot_shock_pct = float(df.iloc[1, 1]) if df.shape[0] > 1 and pd.notna(df.iloc[1, 1]) else 0.0
        vol_shock_abs = float(df.iloc[2, 1]) if df.shape[0] > 2 and pd.notna(df.iloc[2, 1]) else 0.0
        rate_shock_bp = float(df.iloc[3, 1]) if df.shape[0] > 3 and pd.notna(df.iloc[3, 1]) else 0.0
        side = str(df.iloc[4, 1]).strip().upper() if df.shape[0] > 4 and pd.notna(df.iloc[4, 1]) else "COMPRADO"
        qtd_pontos = float(df.iloc[5, 1]) if df.shape[0] > 5 and pd.notna(df.iloc[5, 1]) else 1.0
    except (ValueError, TypeError) as exc:
        scenarios.append(ScenarioInput(
            name=sheet, spot_shock_pct=0.0, vol_shock_abs=0.0,
            rate_shock_bp=0.0, side="COMPRADO", qtd_pontos=1.0,
        ))
        print(f"⚠ Aba '{sheet}' com entradas inválidas: {exc}")
        continue

    scenarios.append(ScenarioInput(
        name=name,
        spot_shock_pct=spot_shock_pct,
        vol_shock_abs=vol_shock_abs,
        rate_shock_bp=rate_shock_bp,
        side=side,
        qtd_pontos=qtd_pontos,
    ))
    print(f"✓ Cenário '{name}': spot={spot_shock_pct:+.1%}, vol={vol_shock_abs:+.4f}, rate={rate_shock_bp:+.0f}bp, lado={side}, qtd={qtd_pontos}")

print(f"\nTotal: {len(scenarios)} cenário(s) carregado(s) de {len(sheet_names)} aba(s).")

# COMMAND ----------

# DBTITLE 1,Execução dos cenários de stress
# Comentário: usa números aleatórios comuns entre base e cenários para reduzir ruído.
from ibov_barrier.stress import run_scenario

rng = np.random.default_rng(sim.seed)
normals = rng.standard_normal((sim.paths, sim.steps))
uniforms = rng.random((sim.paths, sim.steps))

all_results = []
for sc in scenarios:
    result = run_scenario(
        base_market, base_contract, sim, sc,
        normals=normals, uniforms=uniforms,
    )
    all_results.append(result)
    if result.error:
        print(f"⚠ '{result.name}': {result.error}")
    else:
        print(f"✓ '{result.name}': base={result.base_price:,.2f}, estressado={result.stressed_price:,.2f}, P&L={result.pnl:+,.2f} pts{' [KNOCK-OUT]' if result.knock_out else ''}")

# COMMAND ----------

# DBTITLE 1,Geração do STRESS Excel de saída
# Comentário: grava uma nova planilha com os resultados por cenário e uma aba RESUMO.
import openpyxl
from openpyxl.styles import Font, PatternFill

execution_ts = datetime.now(ZoneInfo("America/Sao_Paulo"))
stress_timestamp = execution_ts.strftime("%Y%m%d_%H%M%S")

output_xlsx = results_root / f"STRESS_{stress_timestamp}.xlsx"
output_json = results_root / f"STRESS_{stress_timestamp}.json"

wb = openpyxl.Workbook()
wb.remove(wb.active)

header_font = Font(bold=True, size=11)
input_fill = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
result_fill = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")

INVALID_MSG = "Para estes valores de entrada não temos respaldo matemático para esta simulação."

scenario_lookup = {s.name: s for s in scenarios}

for res in all_results:
    ws = wb.create_sheet(title=res.name[:31])
    sc = scenario_lookup.get(res.name)

    ws["A1"] = "Cenário"
    ws["B1"] = res.name
    ws["A2"] = "Choque de Spot (%)"
    ws["B2"] = sc.spot_shock_pct if sc and not res.error else 0
    ws["A3"] = "Choque de Volatilidade (abs.)"
    ws["B3"] = sc.vol_shock_abs if sc and not res.error else 0
    ws["A4"] = "Choque de Taxa (bp)"
    ws["B4"] = sc.rate_shock_bp if sc and not res.error else 0
    ws["A5"] = "Lado da Posição"
    ws["B5"] = res.side if not res.error else "—"
    ws["A6"] = "QtdPontos"
    ws["B6"] = res.qtd_pontos if not res.error else 0

    for row in range(1, 7):
        ws[f"A{row}"].font = header_font
        ws[f"A{row}"].fill = input_fill

    ws["A9"] = "RESULTADOS"
    ws["A9"].font = Font(bold=True, size=12)

    if res.error:
        ws["A10"] = INVALID_MSG
        ws["A10"].font = Font(italic=True, color="CC0000")
        ws.merge_cells("A10:D10")
    else:
        results_rows = [
            ("Preço-base (pontos)", res.base_price),
            ("Preço estressado (pontos)", res.stressed_price),
            ("Spot estressado", res.spot_stressed),
            ("Volatilidade estressada", res.volatility_stressed),
            ("Taxa estressada", res.rate_stressed),
            ("Knock-out", "SIM" if res.knock_out else "NÃO"),
            ("Lado", res.side),
            ("QtdPontos", res.qtd_pontos),
            ("P&L (pontos)", res.pnl),
        ]
        for i, (label, value) in enumerate(results_rows, start=10):
            ws[f"A{i}"] = label
            ws[f"B{i}"] = value
            ws[f"A{i}"].fill = result_fill
            if isinstance(value, float):
                ws[f"B{i}"].number_format = "#,##0.00"

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 20

# --- Aba RESUMO ---
ws_summary = wb.create_sheet(title="RESUMO")
headers = ["Cenário", "Preço-base", "Preço Estressado", "P&L (pontos)", "Knock-out", "Erro"]
for col_idx, h in enumerate(headers, 1):
    cell = ws_summary.cell(row=1, column=col_idx, value=h)
    cell.font = header_font

for i, res in enumerate(all_results, start=2):
    ws_summary.cell(row=i, column=1, value=res.name)
    ws_summary.cell(row=i, column=2, value=res.base_price if not res.error else "—")
    ws_summary.cell(row=i, column=3, value=res.stressed_price if not res.error else "—")
    ws_summary.cell(row=i, column=4, value=res.pnl if not res.error else "—")
    ws_summary.cell(row=i, column=5, value="SIM" if res.knock_out else ("NÃO" if not res.error else "—"))
    ws_summary.cell(row=i, column=6, value=res.error or "")

for col in "ABCDEF":
    ws_summary.column_dimensions[col].width = 20

wb.save(output_xlsx)
print(f"✓ Excel de saída: {output_xlsx}")

# COMMAND ----------

# DBTITLE 1,Geração do STRESS JSON de auditoria
# Comentário: registra timestamp, arquivos de entrada, hashes, parâmetros, resultados, erros e limitações.
stress_payload = {
    "schema_version": 1,
    "stress_timestamp": execution_ts.isoformat(),
    "input_files": {
        "results_file": results_path.name,
        "results_sha256": results_sha256,
        "validation_file": validation_path.name,
        "risk_file": risk_path.name if risk else None,
        "excel_template": excel_path.name,
    },
    "base_market": {
        "spot": base_market.spot,
        "rate": base_market.rate,
        "dividend_yield": base_market.dividend_yield,
        "volatility": base_market.volatility,
    },
    "base_contract": {
        "strike": base_contract.strike,
        "barrier": base_contract.barrier,
        "maturity": base_contract.maturity,
    },
    "simulation": {
        "paths": sim.paths,
        "steps": sim.steps,
        "seed": sim.seed,
        "monitoring": sim.monitoring,
    },
    "scenarios": [
        {
            "name": r.name,
            "base_price": r.base_price if not r.error else None,
            "stressed_price": r.stressed_price if not r.error else None,
            "pnl": r.pnl if not r.error else None,
            "side": r.side,
            "qtd_pontos": r.qtd_pontos,
            "knock_out": r.knock_out,
            "spot_stressed": r.spot_stressed if not r.error else None,
            "volatility_stressed": r.volatility_stressed if not r.error else None,
            "rate_stressed": r.rate_stressed if not r.error else None,
            "error": r.error,
        }
        for r in all_results
    ],
    "limitations": [
        "P&L expresso em pontos do índice, não em reais (sem multiplicador financeiro validado).",
        "Choques aplicados de forma independente por fator; não modela correlações entre choques.",
        "Reprecificação por Monte Carlo com variável de controle e Brownian Bridge, mesma seed da execução base.",
        "Cenário com spot ≤ barreira resulta em knock-out e preço zero, sem rebate.",
    ],
}

output_json.write_text(json.dumps(stress_payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"✓ JSON de auditoria: {output_json}")

# COMMAND ----------

# DBTITLE 1,Resumo final
print(f"\n{'='*60}")
print(f"STRESS TEST — {execution_ts:%d/%m/%Y %H:%M:%S}")
print(f"{'='*60}")
print(f"Arquivos de entrada:")
print(f"  RESULTS:    {results_path.name}")
print(f"  VALIDATION: {validation_path.name}")
if risk:
    print(f"  RISK:       {risk_path.name}")
print(f"  Excel:      {excel_path.name}")
print(f"\nCenários executados: {len(all_results)}")
for r in all_results:
    status = "ERRO" if r.error else ("KNOCK-OUT" if r.knock_out else "OK")
    pnl_str = f"{r.pnl:+,.2f}" if not r.error else "—"
    print(f"  {r.name:30s} | base={r.base_price:>10,.2f} | stress={r.stressed_price:>10,.2f} | P&L={pnl_str:>12s} | {status}")
print(f"\nSaídas:")
print(f"  {output_xlsx.name}")
print(f"  {output_json.name}")
print(f"\nExcel de entrada preservado: {excel_path}")